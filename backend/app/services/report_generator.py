"""Generación de reportes PDF y Excel para el propietario."""
import io
from datetime import date

from sqlalchemy import func, select

from app.core.database import SessionLocal
from app.models.alerta import Alerta
from app.models.animal import Animal
from app.models.finca import Finca
from app.models.registro_alimentacion import RegistroAlimentacion
from app.models.registro_produccion import RegistroProduccion
from app.models.usuario import Usuario


# ── Almacenamiento en memoria de jobs (MVP) ───────────────────────────────────
jobs: dict[str, dict] = {}


def run_report_job(
    job_id: str,
    finca_id: str,
    fecha_inicio: date,
    fecha_fin: date,
    formato: str,
) -> None:
    """BackgroundTask: genera el reporte y actualiza jobs[job_id]."""
    try:
        data = _collect_data(finca_id, fecha_inicio, fecha_fin)
        if formato == "excel":
            content = _build_excel(data)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"reporte_bovitech_{fecha_inicio}_{fecha_fin}.xlsx"
        else:
            content = _build_pdf(data)
            content_type = "application/pdf"
            filename = f"reporte_bovitech_{fecha_inicio}_{fecha_fin}.pdf"

        jobs[job_id] = {"status": "ready", "data": content, "content_type": content_type, "filename": filename}
    except Exception as exc:
        jobs[job_id] = {"status": "failed", "error": str(exc), "data": None}


def _collect_data(finca_id: str, fecha_inicio: date, fecha_fin: date) -> dict:
    db = SessionLocal()
    try:
        finca = db.get(Finca, finca_id)

        prod_rows = db.execute(
            select(
                Animal.id,
                Animal.numero_arete,
                Animal.nombre,
                func.sum(RegistroProduccion.volumen_litros).label("litros"),
            )
            .join(RegistroProduccion, Animal.id == RegistroProduccion.animal_id)
            .where(
                RegistroProduccion.finca_id == finca_id,
                RegistroProduccion.fecha >= fecha_inicio,
                RegistroProduccion.fecha <= fecha_fin,
            )
            .group_by(Animal.id, Animal.numero_arete, Animal.nombre)
            .order_by(func.sum(RegistroProduccion.volumen_litros).desc())
        ).all()

        feed_rows = db.execute(
            select(
                RegistroAlimentacion.animal_id,
                RegistroAlimentacion.tipo_alimento,
                func.sum(RegistroAlimentacion.cantidad_kg).label("kg"),
            )
            .where(
                RegistroAlimentacion.finca_id == finca_id,
                RegistroAlimentacion.fecha >= fecha_inicio,
                RegistroAlimentacion.fecha <= fecha_fin,
            )
            .group_by(RegistroAlimentacion.animal_id, RegistroAlimentacion.tipo_alimento)
        ).all()

        alertas = db.execute(
            select(Alerta, Animal.numero_arete, Animal.nombre)
            .join(Animal, Alerta.animal_id == Animal.id)
            .where(
                Alerta.finca_id == finca_id,
                Alerta.created_at >= fecha_inicio,
                Alerta.created_at <= fecha_fin,
            )
            .order_by(Alerta.created_at)
        ).all()

        kg_map: dict[str, float] = {}
        for r in feed_rows:
            kg_map[r.animal_id] = kg_map.get(r.animal_id, 0) + float(r.kg)

        animales = []
        total_litros = 0.0
        for r in prod_rows:
            litros = float(r.litros)
            kg = kg_map.get(r.id, 0.0)
            efic = round(litros / kg, 2) if kg > 0 else None
            animales.append({
                "arete": r.numero_arete,
                "nombre": r.nombre or "",
                "litros": round(litros, 1),
                "kg_alimento": round(kg, 1),
                "eficiencia": efic,
            })
            total_litros += litros

        return {
            "finca_nombre": finca.nombre if finca else "Finca",
            "fecha_inicio": str(fecha_inicio),
            "fecha_fin": str(fecha_fin),
            "animales": animales,
            "total_litros": round(total_litros, 1),
            "alertas": [
                {
                    "arete": arete,
                    "nombre": nombre or "",
                    "tipo": a.tipo_alerta,
                    "nivel": a.nivel,
                    "mensaje": a.mensaje,
                    "fecha": str(a.created_at.date()),
                }
                for a, arete, nombre in alertas
            ],
        }
    finally:
        db.close()


def _build_excel(data: dict) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()

    # ── Hoja Producción ───────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Produccion"

    _excel_title(ws, f"BoviTech — Reporte de Producción | {data['finca_nombre']}")
    _excel_subtitle(ws, f"Período: {data['fecha_inicio']} al {data['fecha_fin']}")
    ws.append([])

    headers = ["Arete", "Nombre", "Total Litros", "Kg Alimento", "Eficiencia (L/kg)", "Res. ICA 017/2012"]
    ws.append(headers)
    _style_header_row(ws, ws.max_row)

    for a in data["animales"]:
        ws.append([a["arete"], a["nombre"], a["litros"], a["kg_alimento"], a["eficiencia"] or "—", "✓"])

    ws.append([])
    ws.append(["TOTAL HATO", "", data["total_litros"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 3).font = Font(bold=True)

    _autofit(ws)

    # ── Hoja Alertas ──────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Alertas")
    _excel_title(ws2, "Alertas generadas en el período")
    ws2.append([])
    ws2.append(["Fecha", "Arete", "Nombre", "Tipo", "Nivel", "Mensaje"])
    _style_header_row(ws2, ws2.max_row)
    for al in data["alertas"]:
        ws2.append([al["fecha"], al["arete"], al["nombre"], al["tipo"], al["nivel"], al["mensaje"]])
    _autofit(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pdf(data: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "BoviTech - Reporte de Produccion", ln=True, align="C")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, f"Finca: {data['finca_nombre']}", ln=True, align="C")
    pdf.cell(0, 8, f"Periodo: {data['fecha_inicio']} al {data['fecha_fin']}", ln=True, align="C")
    pdf.ln(6)

    # Tabla producción
    pdf.set_font("Helvetica", "B", 10)
    col_w = [28, 48, 28, 28, 28, 40]
    headers = ["Arete", "Nombre", "Litros", "Kg Alim.", "L/kg", "ICA 017/2012"]
    for i, h in enumerate(headers):
        pdf.cell(col_w[i], 8, h, border=1, align="C")
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for a in data["animales"]:
        row = [a["arete"], a["nombre"][:18], str(a["litros"]), str(a["kg_alimento"]), str(a["eficiencia"] or "—"), "Cumple"]
        for i, val in enumerate(row):
            pdf.cell(col_w[i], 7, val, border=1)
        pdf.ln()

    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(sum(col_w[:3]), 8, f"TOTAL HATO: {data['total_litros']} L", border=1)
    pdf.ln(12)

    # Alertas
    if data["alertas"]:
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 8, "Alertas del periodo", ln=True)
        pdf.set_font("Helvetica", "", 9)
        for al in data["alertas"]:
            pdf.cell(0, 6, f"[{al['nivel'].upper()}] {al['arete']} {al['nombre']} - {al['mensaje']}", ln=True)

    return bytes(pdf.output())


# ── Helpers Excel ─────────────────────────────────────────────────────────────

def _excel_title(ws, text: str) -> None:
    from openpyxl.styles import Font
    ws.append([text])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)


def _excel_subtitle(ws, text: str) -> None:
    ws.append([text])


def _style_header_row(ws, row: int) -> None:
    from openpyxl.styles import Font, PatternFill
    fill = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
